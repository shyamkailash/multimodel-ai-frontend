import tempfile
from pathlib import Path
from django.test import SimpleTestCase

try:
    import pymupdf as fitz
except ImportError:
    import fitz

from pptx import Presentation
from pptx.util import Inches, Pt

from knowledge.extraction.pdf import extract_pdf
from knowledge.extraction.powerpoint import extract_powerpoint
from knowledge.extraction.router import extract_text_from_file


class ExtractionTestCase(SimpleTestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dir_path = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_extract_pdf(self):
        pdf_path = self.dir_path / "sample.pdf"
        doc = fitz.open()

        page1 = doc.new_page()
        page1.insert_text((50, 72), "Machine Learning Basics\nSupervised learning uses labeled data.")

        page2 = doc.new_page()
        page2.insert_text((50, 72), "Deep Learning Overview\nNeural networks have multiple layers.")

        doc.save(str(pdf_path))
        doc.close()

        result = extract_pdf(pdf_path)
        self.assertEqual(result.material_type, "pdf")
        self.assertEqual(len(result.sections), 2)
        self.assertEqual(result.sections[0].number, 1)
        self.assertIn("Supervised learning", result.sections[0].text)
        self.assertEqual(result.sections[1].number, 2)
        self.assertIn("Neural networks", result.sections[1].text)
        self.assertIn("Machine Learning Basics", result.full_text)

    def test_extract_powerpoint(self):
        pptx_path = self.dir_path / "presentation.pptx"
        prs = Presentation()

        slide1 = prs.slides.add_slide(prs.slide_layouts[0])
        title1 = slide1.shapes.title
        title1.text = "Introduction to Computer Vision"
        subtitle1 = slide1.placeholders[1]
        subtitle1.text = "Convolutional Neural Networks"

        slide2 = prs.slides.add_slide(prs.slide_layouts[1])
        title2 = slide2.shapes.title
        title2.text = "Pooling Layers"
        body2 = slide2.placeholders[1]
        body2.text = "Max pooling reduces spatial dimensionality."

        prs.save(str(pptx_path))

        result = extract_powerpoint(pptx_path)
        self.assertEqual(result.material_type, "ppt")
        self.assertEqual(len(result.sections), 2)
        self.assertEqual(result.sections[0].number, 1)
        self.assertIn("Introduction to Computer Vision", result.sections[0].text)
        self.assertEqual(result.sections[1].number, 2)
        self.assertIn("Max pooling", result.sections[1].text)

    def test_extraction_router_pdf(self):
        pdf_path = self.dir_path / "router_test.pdf"
        doc = fitz.open()
        p = doc.new_page()
        p.insert_text((50, 72), "Router test content")
        doc.save(str(pdf_path))
        doc.close()

        result = extract_text_from_file(pdf_path)
        self.assertEqual(result.material_type, "pdf")
        self.assertIn("Router test content", result.full_text)

    def test_extraction_router_unsupported_format(self):
        video_path = self.dir_path / "lecture.mp4"
        video_path.write_text("fake video bytes")

        with self.assertRaises(ValueError) as ctx:
            extract_text_from_file(video_path, material_type="video")
        self.assertIn("Video transcription is not supported", str(ctx.exception))

        unknown_path = self.dir_path / "data.xyz"
        unknown_path.write_text("unknown data")

        with self.assertRaises(ValueError) as ctx:
            extract_text_from_file(unknown_path)
        self.assertIn("Unsupported file type", str(ctx.exception))
